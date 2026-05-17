"""Generate an .xlsx file that mirrors the original 'Penilaian Kinerja Dosen' layout."""
from __future__ import annotations

import io
from collections import defaultdict
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app import models


THIN = Side(border_style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=11)


KD_DESCRIPTIONS = [
    ("KD-1", "Menyampaikan RPS / Kontrak kuliah"),
    ("KD-2", "Kesesuaian materi kuliah dengan RPS / Kontrak Kuliah"),
    ("KD-3", "Kemampuan menjelaskan materi"),
    ("KD-4", "Kemampuan menjawab pertanyaan/ berdiskusi"),
    ("KD-5", "Kemampuan berinteraksi dengan mahasiswa"),
    ("KD-6", "Kedisiplinan / tepat waktu mengajar"),
    ("KD-7", "Kerapian pakaian dosen"),
]


def _write_petunjuk(ws, max_col: int) -> None:
    ws.cell(row=5, column=1, value="Penilaian Mahasiswa pada kinerja Dosen dalam PBM").font = TITLE_FONT
    ws.cell(row=6, column=1, value="SEMESTER GENAP TAHUN AKADEMIK 2025/2026").font = SECTION_FONT
    ws.cell(row=8, column=1, value="Petunjuk:").font = SECTION_FONT
    petunjuk = [
        "1. Berikan penilaian anda terhadap kinerja Dosen dalam Proses Belajar Mengajar (PBM);",
        "2. Lakukan penilaian secara obyektif terhadap setiap unsur yang dinilai;",
        "3. Penilaian anda dijaga kerahasiaannya, dan dijamin aman;",
        "4. Hasil penilaian akan digunakan sebagai dasar dalam upaya perbaikan kinerja Dosen;",
        "5. Gunakan skala 1 - 4 dengan kriteria:",
        "    Nilai 1 : Sangat Kurang",
        "    Nilai 2 : Cukup",
        "    Nilai 3 : Baik",
        "    Nilai 4 : Sangat Baik",
    ]
    for i, txt in enumerate(petunjuk, start=9):
        ws.cell(row=i, column=1, value=txt)

    ws.cell(row=19, column=2, value="KETERANGAN KD").font = SECTION_FONT
    for i, (kode, desc) in enumerate(KD_DESCRIPTIONS, start=20):
        ws.cell(row=i, column=2, value=f"{kode} = {desc}")


def _write_table_header(ws, start_row: int, max_col: int) -> int:
    ws.cell(row=start_row, column=1, value="NO").font = SECTION_FONT
    ws.cell(row=start_row, column=2, value="MATKUL").font = SECTION_FONT
    ws.cell(row=start_row, column=3, value="Nama Dosen").font = SECTION_FONT
    ws.cell(row=start_row, column=4, value="NILAI KINERJA DOSEN (KD)").font = SECTION_FONT
    ws.merge_cells(start_row=start_row, start_column=4, end_row=start_row, end_column=10)
    ws.cell(row=start_row, column=11, value="JML.").font = SECTION_FONT
    ws.cell(row=start_row, column=12, value="RATA-2").font = SECTION_FONT
    ws.cell(row=start_row, column=13, value="SARAN/MASUKAN UNTUK SETIAP DOSEN (JIKA ADA)").font = SECTION_FONT

    sub = start_row + 1
    for i, (kode, _) in enumerate(KD_DESCRIPTIONS, start=4):
        ws.cell(row=sub, column=i, value=kode).font = SECTION_FONT
        ws.cell(row=sub, column=i, value=kode).alignment = Alignment(horizontal="center")

    for col in range(1, max_col + 1):
        for r in (start_row, sub):
            c = ws.cell(row=r, column=col)
            c.fill = HEADER_FILL
            c.border = BORDER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return sub + 1


def build_per_mahasiswa_workbook(
    mahasiswa: models.Mahasiswa,
    penilaian_list: Iterable[models.Penilaian],
) -> bytes:
    """Build an xlsx mirroring the original per-mahasiswa form for a single submission."""
    wb = Workbook()
    ws = wb.active
    ws.title = mahasiswa.prodi.kode if mahasiswa.prodi else "PENILAIAN"
    max_col = 13

    _write_petunjuk(ws, max_col)

    ws.cell(row=28, column=1, value="Nama Mahasiswa").font = SECTION_FONT
    ws.cell(row=28, column=3, value=":")
    ws.cell(row=28, column=4, value=mahasiswa.nama)
    ws.cell(row=29, column=1, value="No. Induk Mahasiswa").font = SECTION_FONT
    ws.cell(row=29, column=3, value=":")
    ws.cell(row=29, column=4, value=mahasiswa.nim)
    ws.cell(row=30, column=1, value="Program Studi").font = SECTION_FONT
    ws.cell(row=30, column=3, value=":")
    ws.cell(row=30, column=4, value=mahasiswa.prodi.nama if mahasiswa.prodi else "")

    data_start = _write_table_header(ws, 32, max_col)

    # Group by matkul: list of (matkul, [penilaian])
    by_matkul: dict[int, list[models.Penilaian]] = defaultdict(list)
    for p in penilaian_list:
        by_matkul[p.kelas.matkul_id].append(p)

    row = data_start
    no = 1
    for matkul_id, plist in by_matkul.items():
        matkul = plist[0].kelas.matkul
        first = True
        for p in plist:
            ws.cell(row=row, column=1, value=no if first else None)
            ws.cell(row=row, column=2, value=f"{matkul.kode} - {matkul.nama}" if first else None)
            ws.cell(row=row, column=3, value=p.kelas.dosen.nama)
            for i, val in enumerate([p.kd1, p.kd2, p.kd3, p.kd4, p.kd5, p.kd6, p.kd7], start=4):
                ws.cell(row=row, column=i, value=val)
            total = p.kd1 + p.kd2 + p.kd3 + p.kd4 + p.kd5 + p.kd6 + p.kd7
            ws.cell(row=row, column=11, value=total)
            ws.cell(row=row, column=12, value=round(total / 7, 2))
            ws.cell(row=row, column=13, value=p.saran or "")
            for c in range(1, max_col + 1):
                ws.cell(row=row, column=c).border = BORDER
                ws.cell(row=row, column=c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            row += 1
            first = False
        no += 1
        row += 1  # blank row between matkul

    widths = [5, 35, 35, 8, 8, 8, 8, 8, 8, 8, 10, 10, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_rekap_workbook(
    periode: models.Periode,
    prodi_groups: dict[str, list[dict]],
) -> bytes:
    """Build a multi-sheet rekap workbook (one sheet per prodi).

    Each row = one dosen-matkul combination with average across all responders.
    """
    wb = Workbook()
    wb.remove(wb.active)

    for prodi_kode, rows in prodi_groups.items():
        ws = wb.create_sheet(title=prodi_kode[:31])
        max_col = 14
        ws.cell(row=1, column=1, value="REKAP PENILAIAN KINERJA DOSEN OLEH MAHASISWA").font = TITLE_FONT
        ws.cell(row=2, column=1, value=f"Program Studi: {rows[0]['prodi_nama'] if rows else prodi_kode}").font = SECTION_FONT
        ws.cell(row=3, column=1, value=f"Semester: {periode.semester} {periode.tahun_akademik}").font = SECTION_FONT

        headers = [
            "NO",
            "Nama Dosen",
            "Matkul",
            "Jml Responden",
            "KD-1",
            "KD-2",
            "KD-3",
            "KD-4",
            "KD-5",
            "KD-6",
            "KD-7",
            "JML",
            "RATA-2",
            "SARAN/MASUKAN",
        ]
        for col, h in enumerate(headers, start=1):
            c = ws.cell(row=5, column=col, value=h)
            c.font = SECTION_FONT
            c.fill = HEADER_FILL
            c.border = BORDER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        row = 6
        for idx, r in enumerate(rows, start=1):
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=r["dosen_nama"])
            ws.cell(row=row, column=3, value=f"{r['matkul_kode']} - {r['matkul_nama']}")
            ws.cell(row=row, column=4, value=r["jumlah_responden"])
            ws.cell(row=row, column=5, value=round(r["avg_kd1"], 2))
            ws.cell(row=row, column=6, value=round(r["avg_kd2"], 2))
            ws.cell(row=row, column=7, value=round(r["avg_kd3"], 2))
            ws.cell(row=row, column=8, value=round(r["avg_kd4"], 2))
            ws.cell(row=row, column=9, value=round(r["avg_kd5"], 2))
            ws.cell(row=row, column=10, value=round(r["avg_kd6"], 2))
            ws.cell(row=row, column=11, value=round(r["avg_kd7"], 2))
            ws.cell(row=row, column=12, value=round(r["jml_total"], 2))
            ws.cell(row=row, column=13, value=round(r["rata_total"], 2))
            ws.cell(row=row, column=14, value="\n".join(r["saran"]))

            for col in range(1, max_col + 1):
                ws.cell(row=row, column=col).border = BORDER
                ws.cell(row=row, column=col).alignment = Alignment(vertical="center", wrap_text=True)
            row += 1

        widths = [5, 30, 35, 12, 8, 8, 8, 8, 8, 8, 8, 8, 10, 40]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    if not wb.sheetnames:
        wb.create_sheet("Rekap")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_raw_workbook(penilaian_rows: list[dict]) -> bytes:
    """Build a flat 'raw responses' workbook for admin."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Raw Responses"

    headers = [
        "Timestamp",
        "NIM",
        "Nama Mahasiswa",
        "Prodi",
        "Matkul (Kode)",
        "Matkul (Nama)",
        "Dosen",
        "KD-1",
        "KD-2",
        "KD-3",
        "KD-4",
        "KD-5",
        "KD-6",
        "KD-7",
        "JML",
        "RATA-2",
        "Saran",
    ]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = SECTION_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, r in enumerate(penilaian_rows, start=2):
        values = [
            r["timestamp"],
            r["nim"],
            r["mhs_nama"],
            r["prodi"],
            r["matkul_kode"],
            r["matkul_nama"],
            r["dosen_nama"],
            r["kd1"],
            r["kd2"],
            r["kd3"],
            r["kd4"],
            r["kd5"],
            r["kd6"],
            r["kd7"],
            r["jml"],
            r["rata"],
            r["saran"],
        ]
        for c_idx, v in enumerate(values, start=1):
            ws.cell(row=r_idx, column=c_idx, value=v).border = BORDER

    widths = [20, 14, 24, 18, 12, 30, 30, 6, 6, 6, 6, 6, 6, 6, 8, 10, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
